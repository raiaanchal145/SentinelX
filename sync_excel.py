"""
SentinelX CLI Excel Synchronization Tool.

Usage:
    python sync_excel.py --sync-only   # Sync test rows from code to Excel without running tests
    python sync_excel.py --dry-run     # Show proposed adds, updates, and removals without changing Excel
"""

import argparse
import ast
import os
import sys
from pathlib import Path
from typing import Any

import openpyxl

from config import EXCEL_PATH, TESTS_DIR
from excel_sync import sync_test_run
from tc import TestCaseMeta, parse_natural_key


def extract_test_cases_from_code(tests_dir: Path = TESTS_DIR) -> dict[str, TestCaseMeta]:
    """
    Parses test files in tests/ using AST to find all @tc decorated functions.
    Does not execute test files, making collection safe and fast.
    """
    collected: dict[str, TestCaseMeta] = {}

    for py_file in sorted(tests_dir.glob("test_*.py")):
        try:
            with open(py_file, "r", encoding="utf-8") as f:
                tree = ast.parse(f.read(), filename=str(py_file))
        except Exception as e:
            print(f"[WARN] Failed to parse {py_file.name}: {e}")
            continue

        for node in ast.walk(tree):
            if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)):
                # Inspect decorators
                for dec in node.decorator_list:
                    # Look for @tc(...)
                    if isinstance(dec, ast.Call) and getattr(dec.func, "id", None) == "tc":
                        kwargs = {}
                        for kw in dec.keywords:
                            if isinstance(kw.value, ast.Constant):
                                kwargs[kw.arg] = kw.value.value

                        tc_id = kwargs.get("id")
                        if tc_id:
                            meta = TestCaseMeta(
                                id=tc_id,
                                module=kwargs.get("module", "General"),
                                description=kwargs.get("description", ""),
                                preconditions=kwargs.get("preconditions", ""),
                                steps=kwargs.get("steps", ""),
                                expected=kwargs.get("expected", ""),
                                type=kwargs.get("type", "automated"),
                                owner=kwargs.get("owner", "Ansh"),
                            )
                            collected[tc_id] = meta

    return collected


def dry_run_sync(collected: dict[str, TestCaseMeta], excel_path: Path = EXCEL_PATH) -> None:
    """Displays differences between test code and Excel report without writing anything."""
    print("\n" + "=" * 70)
    print("SENTINELX EXCEL SYNC - DRY RUN")
    print("=" * 70)

    existing_ids = set()
    if excel_path.exists():
        try:
            wb = openpyxl.load_workbook(excel_path, read_only=True)
            if "Test_Cases" in wb.sheetnames:
                ws = wb["Test_Cases"]
                for row in ws.iter_rows(min_row=2, max_col=1, values_only=True):
                    if row and row[0]:
                        existing_ids.add(str(row[0]).strip())
        except Exception as e:
            print(f"[WARN] Could not read existing Excel file: {e}")

    code_ids = set(collected.keys())

    to_add = sorted(code_ids - existing_ids, key=parse_natural_key)
    to_update = sorted(code_ids & existing_ids, key=parse_natural_key)
    to_remove = sorted(existing_ids - code_ids, key=parse_natural_key)

    print(f"\n[+] NEW CASES TO ADD ({len(to_add)}):")
    if to_add:
        for cid in to_add:
            m = collected[cid]
            print(f"  + [{cid}] ({m.module}) {m.description} [Owner: {m.owner}]")
    else:
        print("  None. All code cases are already in Excel.")

    print(f"\n[*] CASES TO UPDATE ({len(to_update)}):")
    if to_update:
        for cid in to_update:
            m = collected[cid]
            print(f"  * [{cid}] ({m.module}) {m.description}")
    else:
        print("  None.")

    print(f"\n[-] CASES TO REMOVE ({len(to_remove)}):")
    if to_remove:
        for cid in to_remove:
            print(f"  - [{cid}] (No longer found in test code; will be moved to Removed_Cases)")
    else:
        print("  None.")

    print("\n" + "=" * 70)
    print(f"Total in code: {len(code_ids)} | Total in Excel: {len(existing_ids)}")
    print("Dry run completed. No files were modified.")
    print("=" * 70 + "\n")


def main() -> None:
    parser = argparse.ArgumentParser(description="SentinelX Excel Synchronization CLI")
    parser.add_argument("--sync-only", action="store_true", help="Sync test code metadata to Excel without running tests")
    parser.add_argument("--dry-run", action="store_true", help="Show proposed additions, updates, and removals without changing Excel")
    args = parser.parse_args()

    collected = extract_test_cases_from_code()

    if args.dry_run:
        dry_run_sync(collected)
    elif args.sync_only:
        print(f"Syncing {len(collected)} test cases from code to {EXCEL_PATH.name}...")
        saved_file = sync_test_run(
            collected_metas=collected,
            test_results={},
            is_subset=False,
            git_hash="sync-only",
            total_duration=0.0,
        )
        print(f"Excel report synchronized successfully: {saved_file}")
    else:
        parser.print_help()


if __name__ == "__main__":
    main()

