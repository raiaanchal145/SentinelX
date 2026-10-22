"""
SentinelX Automated Excel Report & Test Code Synchronization Engine.

Uses openpyxl to keep test_cases.xlsx strictly in sync with the codebase.
The test code is the single source of truth.
Excel is generated output.
"""

from __future__ import annotations

import datetime as dt
import glob
import os
import shutil
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from config import (
    BACKUP_DIR,
    EXCEL_PATH,
    MAX_BACKUPS,
    SYNC_DELETE_MODE,
)
from tc import TestCaseMeta, parse_natural_key

# Global registry for latest detection metrics passed by test_detection_accuracy
_LATEST_DETECTION_METRICS: dict[str, Any] = {}

def register_detection_metrics(metrics: dict[str, Any]) -> None:
    """Invoked by test_detection_accuracy to supply metrics for the Detection_Metrics sheet."""
    global _LATEST_DETECTION_METRICS
    _LATEST_DETECTION_METRICS = metrics


# ---------------------------------------------------------------------------
# Formatting Styles (Tokens & Colors)
# ---------------------------------------------------------------------------

HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")  # Dark Navy
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)

SECTION_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
SECTION_FILL = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")  # Steel Blue

BOLD_FONT = Font(name="Calibri", size=10, bold=True)
REGULAR_FONT = Font(name="Calibri", size=10, bold=False)

THIN_SIDE = Side(border_style="thin", color="D9D9D9")
CELL_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)

ALIGN_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")

STATUS_STYLES = {
    "PASS": {
        "fill": PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
        "font": Font(name="Calibri", size=10, bold=True, color="006100"),
    },
    "FAIL": {
        "fill": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
        "font": Font(name="Calibri", size=10, bold=True, color="9C0006"),
    },
    "SKIPPED": {
        "fill": PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),
        "font": Font(name="Calibri", size=10, bold=True, color="9C6500"),
    },
    "MANUAL": {
        "fill": PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"),
        "font": Font(name="Calibri", size=10, bold=True, color="B25900"),
    },
    "NOT RUN": {
        "fill": PatternFill(start_color="E7E6E6", end_color="E7E6E6", fill_type="solid"),
        "font": Font(name="Calibri", size=10, bold=False, color="595959"),
    },
}

TEST_CASES_COLUMNS = [
    "TestCaseID",
    "Module",
    "Description",
    "Preconditions",
    "Steps",
    "ExpectedResult",
    "Type",
    "ActualResult",
    "Status",
    "ExecutedOn",
    "Duration(s)",
    "Remarks",
    "Owner",
]


# ---------------------------------------------------------------------------
# Backup & Concurrency Safety
# ---------------------------------------------------------------------------

def create_backup_and_rotate(excel_path: Path) -> str | None:
    """Copies current workbook to backups/ and rotates to keep newest 5."""
    if not excel_path.exists():
        return None

    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    timestamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_file = BACKUP_DIR / f"test_cases_{timestamp}.xlsx"
    shutil.copy2(excel_path, backup_file)

    # Prune old backups
    existing_backups = sorted(
        BACKUP_DIR.glob("test_cases_*.xlsx"),
        key=lambda p: p.stat().st_mtime,
        reverse=True,
    )
    for old_backup in existing_backups[MAX_BACKUPS:]:
        try:
            old_backup.unlink()
        except Exception:
            pass

    return str(backup_file)


def safe_save_workbook(wb: openpyxl.Workbook, target_path: Path) -> Path:
    """
    Saves workbook to target_path. If locked (PermissionError),
    saves to timestamped fallback file and prints a prominent warning.
    """
    try:
        wb.save(target_path)
        return target_path
    except PermissionError:
        timestamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
        fallback_path = target_path.parent / f"test_cases_LOCKED_{timestamp}.xlsx"
        wb.save(fallback_path)
        print("\n" + "*" * 74)
        print("WARNING: 'test_cases.xlsx' is currently locked (likely open in Excel).")
        print(f"Results saved to fallback: {fallback_path.name}")
        print("Please close 'test_cases.xlsx' and run: python sync_excel.py --sync-only")
        print("*" * 74 + "\n")
        return fallback_path


# ---------------------------------------------------------------------------
# Sheet Builders
# ---------------------------------------------------------------------------

def _format_cell(cell, font=REGULAR_FONT, fill=None, alignment=ALIGN_LEFT, border=CELL_BORDER):
    if font:
        cell.font = font
    if fill:
        cell.fill = fill
    if alignment:
        cell.alignment = alignment
    if border:
        cell.border = border


def _autofit_columns(ws, min_width=10, max_width=50):
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or "")
            if "\n" in val_str:
                lines = val_str.split("\n")
                line_len = max(len(l) for l in lines)
            else:
                line_len = len(val_str)
            if line_len > max_len:
                max_len = line_len

        width = max(min_width, min(max_len + 3, max_width))
        ws.column_dimensions[col_letter].width = width


def sync_test_cases_sheet(
    ws,
    collected_metas: dict[str, TestCaseMeta],
    test_results: dict[str, dict[str, Any]],
    is_subset: bool,
    removed_sheet,
) -> dict[str, int]:
    """
    Synchronizes Test_Cases sheet:
    - Preserves existing user Remarks.
    - Updates metadata from code.
    - Updates execution status and actual results.
    - Handles deletions to Removed_Cases on full runs.
    - Applies natural sorting by ID.
    """
    # 1. Read existing rows from Excel
    existing_rows: dict[str, dict[str, Any]] = {}
    if ws.max_row >= 2:
        for row in range(2, ws.max_row + 1):
            tc_id = ws.cell(row=row, column=1).value
            if not tc_id:
                continue
            tc_id_str = str(tc_id).strip()
            existing_rows[tc_id_str] = {
                "TestCaseID": tc_id_str,
                "Module": ws.cell(row=row, column=2).value or "",
                "Description": ws.cell(row=row, column=3).value or "",
                "Preconditions": ws.cell(row=row, column=4).value or "",
                "Steps": ws.cell(row=row, column=5).value or "",
                "ExpectedResult": ws.cell(row=row, column=6).value or "",
                "Type": ws.cell(row=row, column=7).value or "automated",
                "ActualResult": ws.cell(row=row, column=8).value or "",
                "Status": ws.cell(row=row, column=9).value or "NOT RUN",
                "ExecutedOn": ws.cell(row=row, column=10).value or "",
                "Duration(s)": ws.cell(row=row, column=11).value or "",
                "Remarks": ws.cell(row=row, column=12).value or "",
                "Owner": ws.cell(row=row, column=13).value or "",
            }

    # 2. Reconcile with collected test cases from code
    updated_rows: dict[str, dict[str, Any]] = {}
    now_str = dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    # Handle cases present in code
    for tc_id, meta in collected_metas.items():
        res = test_results.get(tc_id)
        if tc_id in existing_rows:
            old = existing_rows[tc_id]
            status = res["status"] if res else old["Status"]
            actual_result = res["actual_result"] if res else old["ActualResult"]
            executed_on = res["executed_on"] if res else old["ExecutedOn"]
            duration = res["duration"] if res else old["Duration(s)"]
            remarks = old.get("Remarks", "")  # NEVER OVERWRITE USER REMARKS
        else:
            status = res["status"] if res else ("MANUAL" if meta.type == "manual" else "NOT RUN")
            actual_result = res["actual_result"] if res else ""
            executed_on = res["executed_on"] if res else ""
            duration = res["duration"] if res else ""
            remarks = ""

        updated_rows[tc_id] = {
            "TestCaseID": tc_id,
            "Module": meta.module,
            "Description": meta.description,
            "Preconditions": meta.preconditions,
            "Steps": meta.steps,
            "ExpectedResult": meta.expected,
            "Type": meta.type,
            "ActualResult": actual_result,
            "Status": status,
            "ExecutedOn": executed_on,
            "Duration(s)": duration,
            "Remarks": remarks,
            "Owner": meta.owner,
        }

    # Handle cases in Excel but removed from code
    for tc_id, old in existing_rows.items():
        if tc_id not in collected_metas:
            if is_subset:
                # Subset run (-k, single file): preserve untouched
                updated_rows[tc_id] = old
            else:
                # Full run: handle deletion according to SYNC_DELETE_MODE
                if SYNC_DELETE_MODE in {"delete", "archive"}:
                    # Append to Removed_Cases sheet
                    rem_row = removed_sheet.max_row + 1
                    removed_sheet.cell(row=rem_row, column=1, value=tc_id)
                    removed_sheet.cell(row=rem_row, column=2, value=old["Module"])
                    removed_sheet.cell(row=rem_row, column=3, value=old["Description"])
                    removed_sheet.cell(row=rem_row, column=4, value=now_str)
                    removed_sheet.cell(row=rem_row, column=5, value=old["Status"])
                    removed_sheet.cell(row=rem_row, column=6, value="Removed from test code")
                    for col_idx in range(1, 7):
                        _format_cell(removed_sheet.cell(row=rem_row, column=col_idx))
                elif SYNC_DELETE_MODE == "keep_flagged":
                    old["Remarks"] = f"[FLAGGED REMOVED {now_str}] {old.get('Remarks', '')}".strip()
                    updated_rows[tc_id] = old

    # 3. Sort rows naturally (AUTH-02 before AUTH-10)
    sorted_ids = sorted(updated_rows.keys(), key=parse_natural_key)

    # 4. Clear and rebuild Test_Cases sheet
    ws.delete_rows(1, ws.max_row + 1)

    # Header
    ws.append(TEST_CASES_COLUMNS)
    ws.row_dimensions[1].height = 28
    for col_idx in range(1, len(TEST_CASES_COLUMNS) + 1):
        cell = ws.cell(row=1, column=col_idx)
        _format_cell(cell, font=HEADER_FONT, fill=HEADER_FILL, alignment=HEADER_ALIGN)

    # Data Rows
    status_counts = {"PASS": 0, "FAIL": 0, "SKIPPED": 0, "MANUAL": 0, "NOT RUN": 0}
    for row_idx, tc_id in enumerate(sorted_ids, start=2):
        row_data = updated_rows[tc_id]
        status = row_data["Status"]
        status_counts[status] = status_counts.get(status, 0) + 1

        values = [row_data[col] for col in TEST_CASES_COLUMNS]
        ws.append(values)
        ws.row_dimensions[row_idx].height = 20

        # Format row cells
        for col_idx in range(1, len(TEST_CASES_COLUMNS) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            if col_idx == 1:  # ID
                _format_cell(cell, font=BOLD_FONT, alignment=ALIGN_CENTER)
            elif col_idx in {7, 10, 11, 13}:  # Type, ExecutedOn, Duration, Owner
                _format_cell(cell, alignment=ALIGN_CENTER)
            elif col_idx == 9:  # Status
                style = STATUS_STYLES.get(status, STATUS_STYLES["NOT RUN"])
                _format_cell(cell, font=style["font"], fill=style["fill"], alignment=ALIGN_CENTER)
            else:
                _format_cell(cell, alignment=ALIGN_LEFT)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(TEST_CASES_COLUMNS))}{ws.max_row}"
    _autofit_columns(ws)

    return status_counts


def sync_run_history_sheet(
    ws,
    status_counts: dict[str, int],
    git_hash: str,
    total_duration: float,
) -> None:
    """Appends execution record row to Run_History."""
    if ws.max_row == 1 and ws.cell(row=1, column=1).value is None:
        ws.delete_rows(1, 1)

    if ws.max_row == 0 or ws.cell(row=1, column=1).value != "Run Timestamp":
        headers = [
            "Run Timestamp",
            "Total Tests",
            "Passed",
            "Failed",
            "Skipped",
            "Manual",
            "Pass Rate (%)",
            "Git Hash",
            "Duration (s)",
        ]
        ws.append(headers)
        ws.row_dimensions[1].height = 26
        for c in range(1, len(headers) + 1):
            _format_cell(ws.cell(row=1, column=c), font=HEADER_FONT, fill=HEADER_FILL, alignment=HEADER_ALIGN)

    total = sum(status_counts.values())
    passed = status_counts.get("PASS", 0)
    failed = status_counts.get("FAIL", 0)
    skipped = status_counts.get("SKIPPED", 0)
    manual = status_counts.get("MANUAL", 0)
    pass_rate = round((passed / total * 100), 1) if total > 0 else 0.0

    now_str = dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    row_values = [
        now_str,
        total,
        passed,
        failed,
        skipped,
        manual,
        f"{pass_rate}%",
        git_hash,
        total_duration,
    ]
    ws.append(row_values)
    new_row = ws.max_row
    ws.row_dimensions[new_row].height = 20
    for c in range(1, len(row_values) + 1):
        _format_cell(ws.cell(row=new_row, column=c), alignment=ALIGN_CENTER)

    ws.freeze_panes = "A2"
    _autofit_columns(ws)


def sync_summary_sheet(ws, tc_sheet) -> None:
    """Rebuilds per-module test counts and pass rates on Summary sheet."""
    ws.delete_rows(1, ws.max_row + 1)

    headers = ["Module", "Total", "Pass", "Fail", "Skipped", "Manual", "Not Run", "Pass Rate (%)"]
    ws.append(headers)
    ws.row_dimensions[1].height = 26
    for c in range(1, len(headers) + 1):
        _format_cell(ws.cell(row=1, column=c), font=HEADER_FONT, fill=HEADER_FILL, alignment=HEADER_ALIGN)

    # Aggregate by module from tc_sheet
    by_module: dict[str, dict[str, int]] = {}
    for r in range(2, tc_sheet.max_row + 1):
        mod = str(tc_sheet.cell(row=r, column=2).value or "Other").strip()
        status = str(tc_sheet.cell(row=r, column=9).value or "NOT RUN").strip()
        if mod not in by_module:
            by_module[mod] = {"Total": 0, "PASS": 0, "FAIL": 0, "SKIPPED": 0, "MANUAL": 0, "NOT RUN": 0}
        by_module[mod]["Total"] += 1
        by_module[mod][status] = by_module[mod].get(status, 0) + 1

    totals = {"Total": 0, "PASS": 0, "FAIL": 0, "SKIPPED": 0, "MANUAL": 0, "NOT RUN": 0}
    for mod in sorted(by_module.keys()):
        stats = by_module[mod]
        for k in totals:
            totals[k] += stats[k]

        pass_rate = round((stats["PASS"] / stats["Total"] * 100), 1) if stats["Total"] > 0 else 0.0
        row_vals = [
            mod,
            stats["Total"],
            stats["PASS"],
            stats["FAIL"],
            stats["SKIPPED"],
            stats["MANUAL"],
            stats["NOT RUN"],
            f"{pass_rate}%",
        ]
        ws.append(row_vals)
        cur_row = ws.max_row
        ws.row_dimensions[cur_row].height = 20
        _format_cell(ws.cell(row=cur_row, column=1), font=BOLD_FONT, alignment=ALIGN_LEFT)
        for c in range(2, len(row_vals) + 1):
            _format_cell(ws.cell(row=cur_row, column=c), alignment=ALIGN_CENTER)

    # Overall Totals Row
    total_pass_rate = round((totals["PASS"] / totals["Total"] * 100), 1) if totals["Total"] > 0 else 0.0
    tot_row_vals = [
        "TOTAL",
        totals["Total"],
        totals["PASS"],
        totals["FAIL"],
        totals["SKIPPED"],
        totals["MANUAL"],
        totals["NOT RUN"],
        f"{total_pass_rate}%",
    ]
    ws.append(tot_row_vals)
    tot_row = ws.max_row
    ws.row_dimensions[tot_row].height = 22
    for c in range(1, len(tot_row_vals) + 1):
        _format_cell(ws.cell(row=tot_row, column=c), font=BOLD_FONT, fill=SECTION_FILL, alignment=ALIGN_CENTER)

    ws.freeze_panes = "A2"
    _autofit_columns(ws)


def sync_detection_metrics_sheet(ws, metrics: dict[str, Any]) -> None:
    """Populates Detection_Metrics sheet with confusion matrix and severity tables."""
    ws.delete_rows(1, ws.max_row + 1)

    # Title Banner
    ws.cell(row=1, column=1, value="SentinelX AI Detection Accuracy & Confusion Matrix")
    ws.merge_cells("A1:E1")
    _format_cell(ws.cell(row=1, column=1), font=HEADER_FONT, fill=HEADER_FILL, alignment=ALIGN_CENTER)
    ws.row_dimensions[1].height = 28

    if not metrics:
        # Load from tests/data/manifest.json if empty
        try:
            from tests.test_detection_accuracy import compute_dataset_metrics
            metrics = compute_dataset_metrics()
        except Exception:
            metrics = {}

    cm = metrics.get("confusion_matrix", {"TP": 0, "FP": 0, "TN": 0, "FN": 0})
    m = metrics.get("metrics", {"accuracy": 0, "precision": 0, "recall": 0, "f1_score": 0, "threshold": 0.80})

    # Section 1: Confusion Matrix Table
    ws.cell(row=3, column=1, value="Classification Outcome")
    ws.cell(row=3, column=2, value="Count")
    ws.cell(row=3, column=3, value="Performance Metric")
    ws.cell(row=3, column=4, value="Score")
    ws.cell(row=3, column=5, value="Target Threshold")

    ws.row_dimensions[3].height = 24
    for c in range(1, 6):
        _format_cell(ws.cell(row=3, column=c), font=BOLD_FONT, fill=SECTION_FILL, alignment=ALIGN_CENTER)

    matrix_rows = [
        ("True Positives (TP)", cm["TP"], "Accuracy", f"{m['accuracy'] * 100:.1f}%", f">= {m['threshold'] * 100:.0f}%"),
        ("False Positives (FP)", cm["FP"], "Precision", f"{m['precision'] * 100:.1f}%", f">= {m['threshold'] * 100:.0f}%"),
        ("True Negatives (TN)", cm["TN"], "Recall", f"{m['recall'] * 100:.1f}%", f">= {m['threshold'] * 100:.0f}%"),
        ("False Negatives (FN)", cm["FN"], "F1-Score", f"{m['f1_score'] * 100:.1f}%", f">= {m['threshold'] * 100:.0f}%"),
    ]

    for idx, rdata in enumerate(matrix_rows, start=4):
        ws.row_dimensions[idx].height = 20
        ws.cell(row=idx, column=1, value=rdata[0])
        ws.cell(row=idx, column=2, value=rdata[1])
        ws.cell(row=idx, column=3, value=rdata[2])
        ws.cell(row=idx, column=4, value=rdata[3])
        ws.cell(row=idx, column=5, value=rdata[4])

        _format_cell(ws.cell(row=idx, column=1), alignment=ALIGN_LEFT)
        _format_cell(ws.cell(row=idx, column=2), alignment=ALIGN_CENTER)
        _format_cell(ws.cell(row=idx, column=3), font=BOLD_FONT, alignment=ALIGN_LEFT)
        _format_cell(ws.cell(row=idx, column=4), font=BOLD_FONT, alignment=ALIGN_CENTER)
        _format_cell(ws.cell(row=idx, column=5), alignment=ALIGN_CENTER)

    # Threshold verdict banner
    met = m.get("threshold_met", True)
    verdict_text = "VERDICT: ACCURACY THRESHOLD TARGET MET (PASS)" if met else "VERDICT: BELOW REQUIRED THRESHOLD (FAIL)"
    verdict_fill = STATUS_STYLES["PASS"]["fill"] if met else STATUS_STYLES["FAIL"]["fill"]
    verdict_font = STATUS_STYLES["PASS"]["font"] if met else STATUS_STYLES["FAIL"]["font"]

    ws.cell(row=8, column=1, value=verdict_text)
    ws.merge_cells("A8:E8")
    ws.row_dimensions[8].height = 24
    _format_cell(ws.cell(row=8, column=1), font=verdict_font, fill=verdict_fill, alignment=ALIGN_CENTER)

    # Section 2: Per-Severity Breakdown
    ws.cell(row=10, column=1, value="Per-Severity Detection Breakdown")
    ws.merge_cells("A10:E10")
    _format_cell(ws.cell(row=10, column=1), font=BOLD_FONT, fill=SECTION_FILL, alignment=ALIGN_CENTER)
    ws.row_dimensions[10].height = 24

    sev_headers = ["Severity", "Expected Malicious", "Detected Correctly", "Missed (FN)", "Detection Rate (%)"]
    ws.append(sev_headers)
    ws.row_dimensions[11].height = 22
    for c in range(1, 6):
        _format_cell(ws.cell(row=11, column=c), font=BOLD_FONT, fill=SECTION_FILL, alignment=ALIGN_CENTER)

    per_sev = metrics.get("per_severity", {})
    for idx, sev in enumerate(["Critical", "High", "Medium", "Low"], start=12):
        s_data = per_sev.get(sev, {"expected": 0, "detected": 0, "missed": 0})
        exp = s_data["expected"]
        det = s_data["detected"]
        mis = s_data["missed"]
        rate = f"{(det / exp * 100):.1f}%" if exp > 0 else "N/A"

        ws.cell(row=idx, column=1, value=sev)
        ws.cell(row=idx, column=2, value=exp)
        ws.cell(row=idx, column=3, value=det)
        ws.cell(row=idx, column=4, value=mis)
        ws.cell(row=idx, column=5, value=rate)
        ws.row_dimensions[idx].height = 20

        _format_cell(ws.cell(row=idx, column=1), font=BOLD_FONT, alignment=ALIGN_LEFT)
        for c in range(2, 6):
            _format_cell(ws.cell(row=idx, column=c), alignment=ALIGN_CENTER)

    _autofit_columns(ws)


def sync_removed_cases_sheet(ws) -> None:
    """Initializes Removed_Cases headers if sheet is empty."""
    headers = ["TestCaseID", "Module", "Description", "RemovedAt", "LastStatus", "Reason"]
    if ws.max_row == 1 and ws.cell(row=1, column=1).value is None:
        ws.delete_rows(1, 1)

    if ws.max_row == 0 or ws.cell(row=1, column=1).value != "TestCaseID":
        ws.delete_rows(1, ws.max_row + 1)
        ws.append(headers)
        ws.row_dimensions[1].height = 26
        for c in range(1, len(headers) + 1):
            _format_cell(ws.cell(row=1, column=c), font=HEADER_FONT, fill=HEADER_FILL, alignment=HEADER_ALIGN)
        ws.freeze_panes = "A2"
        _autofit_columns(ws)



# ---------------------------------------------------------------------------
# Main Orchestrator
# ---------------------------------------------------------------------------

def sync_test_run(
    collected_metas: dict[str, TestCaseMeta],
    test_results: dict[str, dict[str, Any]],
    is_subset: bool = False,
    git_hash: str = "unknown",
    total_duration: float = 0.0,
    excel_path: Path = EXCEL_PATH,
) -> Path:
    """
    Main entrypoint called by pytest_sessionfinish or CLI sync.
    Creates or updates test_cases.xlsx with full formatting and backup rotation.
    """
    # 1. Backup existing workbook if present
    create_backup_and_rotate(excel_path)

    # 2. Open or create workbook
    if excel_path.exists():
        try:
            wb = openpyxl.load_workbook(excel_path)
        except Exception:
            wb = openpyxl.Workbook()
    else:
        wb = openpyxl.Workbook()

    # Ensure required sheets exist
    sheet_names = wb.sheetnames
    if "Test_Cases" not in sheet_names:
        ws_tc = wb.active
        ws_tc.title = "Test_Cases"
    else:
        ws_tc = wb["Test_Cases"]

    ws_history = wb["Run_History"] if "Run_History" in sheet_names else wb.create_sheet("Run_History")
    ws_summary = wb["Summary"] if "Summary" in sheet_names else wb.create_sheet("Summary")
    ws_detection = wb["Detection_Metrics"] if "Detection_Metrics" in sheet_names else wb.create_sheet("Detection_Metrics")
    ws_removed = wb["Removed_Cases"] if "Removed_Cases" in sheet_names else wb.create_sheet("Removed_Cases")

    # Initialize Removed_Cases headers
    sync_removed_cases_sheet(ws_removed)

    # Sync Test_Cases
    status_counts = sync_test_cases_sheet(
        ws_tc,
        collected_metas=collected_metas,
        test_results=test_results,
        is_subset=is_subset,
        removed_sheet=ws_removed,
    )

    # Sync Run_History
    sync_run_history_sheet(
        ws_history,
        status_counts=status_counts,
        git_hash=git_hash,
        total_duration=total_duration,
    )

    # Sync Summary
    sync_summary_sheet(ws_summary, ws_tc)

    # Sync Detection_Metrics
    sync_detection_metrics_sheet(ws_detection, _LATEST_DETECTION_METRICS)

    # Save safely (with locked file fallback if needed)
    saved_path = safe_save_workbook(wb, excel_path)
    return saved_path
