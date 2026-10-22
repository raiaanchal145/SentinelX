"""
SentinelX QA Test Suite: Excel Synchronization Engine & Safety Logic.

Module: SYNC
Owners: Ansh, Riteka Singh, Aanchal Rai
"""

import os
from pathlib import Path
from unittest.mock import patch
import openpyxl
import pytest

from config import validate_test_database_url
from excel_sync import sync_test_run
from tc import TestCaseMeta, tc


@tc(
    id="SYNC-01",
    module="SyncLogic",
    description="Sync creates test_cases.xlsx with standard sheets when file does not exist",
    preconditions="Target Excel file does not exist",
    steps="Invoke sync_test_run on non-existent path",
    expected="Workbook created with Test_Cases, Run_History, Summary, Detection_Metrics, Removed_Cases sheets",
    type="automated",
    owner="Ansh",
)
def test_sync_01_create_missing_file(tmp_path: Path):
    target_excel = tmp_path / "test_cases.xlsx"
    assert not target_excel.exists()

    meta = TestCaseMeta(
        id="AUTH-01",
        module="Auth",
        description="Login test",
        preconditions="User exists",
        steps="Post credentials",
        expected="Token returned",
        type="automated",
        owner="Ansh",
    )

    saved_path = sync_test_run(
        collected_metas={"AUTH-01": meta},
        test_results={},
        is_subset=False,
        excel_path=target_excel,
    )

    assert target_excel.exists()
    wb = openpyxl.load_workbook(target_excel)
    assert "Test_Cases" in wb.sheetnames
    assert "Run_History" in wb.sheetnames
    assert "Summary" in wb.sheetnames
    assert "Detection_Metrics" in wb.sheetnames
    assert "Removed_Cases" in wb.sheetnames


@tc(
    id="SYNC-02",
    module="SyncLogic",
    description="New test in code adds new row with Status 'NOT RUN' when no results provided",
    preconditions="Existing Excel file with 1 test case",
    steps="Sync with an additional test case in code",
    expected="Test_Cases sheet has 2 rows and new test has status 'NOT RUN'",
    type="automated",
    owner="Riteka Singh",
)
def test_sync_02_add_new_test_row(tmp_path: Path):
    target_excel = tmp_path / "test_cases.xlsx"

    meta1 = TestCaseMeta(
        id="AUTH-01",
        module="Auth",
        description="First test",
        preconditions="Pre 1",
        steps="Step 1",
        expected="Exp 1",
        type="automated",
        owner="Ansh",
    )
    meta2 = TestCaseMeta(
        id="AUTH-02",
        module="Auth",
        description="Second test",
        preconditions="Pre 2",
        steps="Step 2",
        expected="Exp 2",
        type="automated",
        owner="Riteka Singh",
    )

    # Initial run with 1 case
    sync_test_run(
        collected_metas={"AUTH-01": meta1},
        test_results={"AUTH-01": {"status": "PASS", "actual_result": "Success", "duration": 0.1, "executed_on": "now"}},
        excel_path=target_excel,
    )

    # Second run with 2 cases (AUTH-02 has no result yet)
    sync_test_run(
        collected_metas={"AUTH-01": meta1, "AUTH-02": meta2},
        test_results={"AUTH-01": {"status": "PASS", "actual_result": "Success", "duration": 0.1, "executed_on": "now"}},
        excel_path=target_excel,
    )

    wb = openpyxl.load_workbook(target_excel)
    ws = wb["Test_Cases"]
    assert ws.max_row == 3  # Header + 2 rows
    assert ws.cell(row=2, column=1).value == "AUTH-01"
    assert ws.cell(row=3, column=1).value == "AUTH-02"
    assert ws.cell(row=3, column=9).value == "NOT RUN"


@tc(
    id="SYNC-03",
    module="SyncLogic",
    description="Test removed from code is moved from Test_Cases to Removed_Cases sheet",
    preconditions="Existing Excel file contains AUTH-01 and AUTH-02",
    steps="Perform full sync with only AUTH-01 in code",
    expected="AUTH-02 deleted from Test_Cases and recorded in Removed_Cases with reason",
    type="automated",
    owner="Aanchal Rai",
)
def test_sync_03_remove_test_row(tmp_path: Path):
    target_excel = tmp_path / "test_cases.xlsx"

    meta1 = TestCaseMeta(id="AUTH-01", module="Auth", description="Test 1", preconditions="", steps="", expected="", type="automated", owner="Ansh")
    meta2 = TestCaseMeta(id="AUTH-02", module="Auth", description="Test 2", preconditions="", steps="", expected="", type="automated", owner="Riteka Singh")

    sync_test_run(
        collected_metas={"AUTH-01": meta1, "AUTH-02": meta2},
        test_results={"AUTH-02": {"status": "PASS", "actual_result": "Passed", "duration": 0.2, "executed_on": "now"}},
        excel_path=target_excel,
    )

    # Sync with AUTH-02 omitted (removed from code)
    sync_test_run(
        collected_metas={"AUTH-01": meta1},
        test_results={"AUTH-01": {"status": "PASS", "actual_result": "Passed", "duration": 0.1, "executed_on": "now"}},
        is_subset=False,
        excel_path=target_excel,
    )

    wb = openpyxl.load_workbook(target_excel)
    ws_tc = wb["Test_Cases"]
    assert ws_tc.max_row == 2  # Header + AUTH-01 only
    assert ws_tc.cell(row=2, column=1).value == "AUTH-01"

    ws_rem = wb["Removed_Cases"]
    assert ws_rem.max_row == 2  # Header + AUTH-02
    assert ws_rem.cell(row=2, column=1).value == "AUTH-02"
    assert ws_rem.cell(row=2, column=6).value == "Removed from test code"


@tc(
    id="SYNC-04",
    module="SyncLogic",
    description="User remarks edited directly in Excel are strictly preserved on subsequent sync runs",
    preconditions="User has typed notes into Remarks column in Excel",
    steps="Re-run sync from test code",
    expected="Remarks cell content remains intact and is never overwritten",
    type="automated",
    owner="Ansh",
)
def test_sync_04_preserve_user_remarks(tmp_path: Path):
    target_excel = tmp_path / "test_cases.xlsx"

    meta1 = TestCaseMeta(id="AUTH-01", module="Auth", description="Original Desc", preconditions="", steps="", expected="", type="automated", owner="Ansh")

    sync_test_run(
        collected_metas={"AUTH-01": meta1},
        test_results={"AUTH-01": {"status": "PASS", "actual_result": "OK", "duration": 0.1, "executed_on": "now"}},
        excel_path=target_excel,
    )

    # User manually edits Remarks column (column 12)
    wb = openpyxl.load_workbook(target_excel)
    ws = wb["Test_Cases"]
    ws.cell(row=2, column=12, value="Manual QA note: verified on staging cluster")
    wb.save(target_excel)

    # Code description changes and re-syncs
    meta1_updated = TestCaseMeta(id="AUTH-01", module="Auth", description="Updated Description from Code", preconditions="", steps="", expected="", type="automated", owner="Ansh")

    sync_test_run(
        collected_metas={"AUTH-01": meta1_updated},
        test_results={"AUTH-01": {"status": "PASS", "actual_result": "OK", "duration": 0.1, "executed_on": "now"}},
        excel_path=target_excel,
    )

    # Verify remarks were preserved while description updated
    wb2 = openpyxl.load_workbook(target_excel)
    ws2 = wb2["Test_Cases"]
    assert ws2.cell(row=2, column=3).value == "Updated Description from Code"
    assert ws2.cell(row=2, column=12).value == "Manual QA note: verified on staging cluster"


@tc(
    id="SYNC-05",
    module="SyncLogic",
    description="Locked Excel file (PermissionError) saves timestamped fallback copy without crashing",
    preconditions="Target Excel file raises PermissionError on save",
    steps="Invoke sync_test_run while target path raises PermissionError",
    expected="Fallback file test_cases_LOCKED_*.xlsx created and function returns cleanly",
    type="automated",
    owner="Riteka Singh",
)
def test_sync_05_handle_locked_file(tmp_path: Path):
    target_excel = tmp_path / "test_cases.xlsx"
    meta1 = TestCaseMeta(id="AUTH-01", module="Auth", description="Desc", preconditions="", steps="", expected="", type="automated", owner="Ansh")

    # Simulate PermissionError on save of target_excel
    original_save = openpyxl.Workbook.save

    def mock_save(self, filename):
        if str(filename) == str(target_excel):
            raise PermissionError("File is locked by another process.")
        return original_save(self, filename)

    with patch.object(openpyxl.Workbook, "save", side_effect=mock_save, autospec=True):
        saved_path = sync_test_run(
            collected_metas={"AUTH-01": meta1},
            test_results={},
            excel_path=target_excel,
        )

    assert "test_cases_LOCKED_" in str(saved_path)
    assert saved_path.exists()


@tc(
    id="SYNC-06",
    module="SyncLogic",
    description="Database safety guard strictly refuses shared or non-test databases",
    preconditions="Database URL contains cloud hosted provider or lacks 'test'",
    steps="Call validate_test_database_url with invalid DB URLs",
    expected="RuntimeError raised with explicit safety warning",
    type="automated",
    owner="Aanchal Rai",
)
def test_sync_06_database_safety_guard():
    with pytest.raises(RuntimeError):
        validate_test_database_url("postgresql+asyncpg://admin:pass@db.supabase.co:5432/live")

    with pytest.raises(RuntimeError):
        validate_test_database_url("postgresql+asyncpg://localhost:5432/production_data")

    with pytest.raises(RuntimeError):
        validate_test_database_url("")
